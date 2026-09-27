"""Full CLI run against fake BMCs: CSV in, servers.csv / firmware.csv / report.xlsx out, then resume."""

import csv
from pathlib import Path

import pytest
from openpyxl import load_workbook

import fwtool.redfish.client as client_mod
from fakebmc import RouterSession, conn_error
from fwtool.cli import main

FX = Path(__file__).parent / "fixtures"

FLEET = {
    "10.0.0.1": "idrac9_r740",
    "10.0.0.2": "idrac9_vxrail",
    "10.0.0.3": "idrac8_r630",
    "10.0.0.4": "idrac8_partial",
    "10.0.0.5": "ilo5_dl380g10",
    "10.0.0.6": "ilo6_dl380g11",
    "10.0.0.7": "ilo4_dl380g9",
    "10.0.0.8": conn_error(),
}


@pytest.fixture
def fake_fleet(monkeypatch):
    router = RouterSession(FLEET)
    monkeypatch.setattr(client_mod, "_make_session", lambda: router)
    monkeypatch.setenv("DELL_BMC_USER", "ro")
    monkeypatch.setenv("DELL_BMC_PASS", "secret")
    monkeypatch.setenv("HPE_BMC_USER", "ro")
    monkeypatch.setenv("HPE_BMC_PASS", "secret")
    monkeypatch.delenv("BMC_USER", raising=False)
    return router


def write_input(tmp_path):
    p = tmp_path / "servers.csv"
    rows = ["﻿name,bmc_ip,support,site,environment", ""]
    for i, ip in enumerate(FLEET, 1):
        support = "tpm" if ip == "10.0.0.6" else ("hpe" if ip in ("10.0.0.5", "10.0.0.7") else "")
        rows.append(f"srv{i:02d},{ip},{support},LON,{'prod' if i % 2 else 'nonprod'}")
    rows.append("dup,10.0.0.1,,,")
    p.write_text("\n".join(rows) + "\n", encoding="utf-8")
    return p


def args(tmp_path, *extra):
    return ["inventory", "-i", str(write_input(tmp_path)), "-o", str(tmp_path / "runs"), "--no-prompt",
            "--no-progress", "--retries", "0", "--workers", "4", "--offline",
            "--dell-catalog", str(FX / "dell_catalog.xml"), "--hpe-reference", str(FX / "hpe_reference.yaml"),
            "--baselines", str(FX / "baselines.yaml"), "--as-of", "2026-09-27", *extra]


def read_csv(p):
    with open(p, encoding="utf-8-sig") as fh:
        return list(csv.DictReader(fh))


def test_full_run_and_resume(tmp_path, fake_fleet, capsys):
    assert main(args(tmp_path, "--save-raw")) == 0
    run_dir = next((tmp_path / "runs").iterdir())
    servers = read_csv(run_dir / "servers.csv")
    assert len(servers) == 8
    by_ip = {s["bmc_ip"]: s for s in servers}
    assert by_ip["10.0.0.2"]["vxrail"] == "yes" and by_ip["10.0.0.2"]["update_eligible"].startswith("no")
    assert by_ip["10.0.0.4"]["collection_status"] == "partial"
    assert by_ip["10.0.0.7"]["collection_status"] == "partial"
    assert by_ip["10.0.0.8"]["collection_status"] == "failed" and by_ip["10.0.0.8"]["error_class"] == "unreachable"
    assert by_ip["10.0.0.8"]["fix_hint"]
    assert by_ip["10.0.0.6"]["support"] == "tpm"
    # ranked worst first; failed servers last and unranked
    ranks = [s["rank"] for s in servers]
    assert ranks[:7] == [str(i) for i in range(1, 8)] and ranks[7] == ""
    scores = [float(s["priority_score"]) for s in servers[:7]]
    assert scores == sorted(scores, reverse=True)

    fw = read_csv(run_dir / "firmware.csv")
    assert {f["category"] for f in fw} == {"BIOS", "BMC", "CPLD", "Storage", "NIC", "Drive", "PSU", "Other"}
    assert any(f["update_track"] == "entitlement required" for f in fw)

    wb = load_workbook(run_dir / "report.xlsx", read_only=True)
    assert wb.sheetnames[:4] == ["Summary", "Worst 100", "Model x Component", "Failures & Partials"]
    fails = list(wb["Failures & Partials"].iter_rows(values_only=True))
    assert len(fails) == 1 + 3  # header + unreachable + 2 partial

    assert (run_dir / "raw" / "10.0.0.5.json").exists()
    # credentials never written anywhere
    for p in run_dir.rglob("*"):
        if p.is_file() and p.suffix in (".json", ".csv", ".log"):
            assert "secret" not in p.read_text(encoding="utf-8", errors="ignore"), p

    # resume: collected servers are skipped, the failed one is retried
    before = len(fake_fleet.hosts)
    capsys.readouterr()
    assert main(args(tmp_path, "--resume", str(run_dir))) == 0
    out = capsys.readouterr().out
    assert "skipped(resume)=7" in out and "failed=1" in out
    assert set(fake_fleet.hosts[before:]) == {"10.0.0.8"}  # only the failed host was contacted
    assert len(read_csv(run_dir / "servers.csv")) == 8


def test_report_regeneration(tmp_path, fake_fleet):
    assert main(args(tmp_path)) == 0
    run_dir = next((tmp_path / "runs").iterdir())
    (run_dir / "servers.csv").unlink()
    assert main(["report", "--run-dir", str(run_dir), "--offline", "--dell-catalog", str(FX / "dell_catalog.xml"),
                 "--hpe-reference", str(FX / "hpe_reference.yaml"), "--target", "latest"]) == 0
    assert (run_dir / "servers.csv").exists()


def test_pilot_limit(tmp_path, fake_fleet):
    assert main(args(tmp_path, "--limit", "2")) == 0
    run_dir = next((tmp_path / "runs").iterdir())
    assert len(read_csv(run_dir / "servers.csv")) == 2


def test_validate_command(tmp_path, fake_fleet, capsys, monkeypatch):
    assert main(args(tmp_path)) == 0
    monkeypatch.chdir(tmp_path)
    capsys.readouterr()
    ref = ["--dell-catalog", str(FX / "dell_catalog.xml"), "--hpe-reference", str(FX / "hpe_reference.yaml")]
    rc = main(["validate", *ref])
    out = capsys.readouterr().out
    assert rc == 0  # warnings (unreachable, partial) but no hard failure
    assert "[WARN] Collection: 8 servers: ok=5 partial=2 failed=1" in out
    assert "[WARN] Failed: unreachable" in out and "10.0.0.8" in out
    assert "[PASS] BIOS and BMC firmware found" in out
    assert "[PASS] HPE BIOS version parsed" in out
    assert "HPE reference coverage" in out and "Appliance nodes flagged" in out
    run_dir = next((tmp_path / "runs").iterdir())
    assert (run_dir / "validation.txt").exists()
    assert main(["validate", "--strict", *ref]) == 1

    capsys.readouterr()
    main(["validate", "--share", *ref])
    shared = (run_dir / "validation-share.txt").read_text(encoding="utf-8")
    assert "10.0.0." not in shared and "srv01" not in shared
    assert "server001" in shared and "ip001" in shared
    assert "secret" not in shared


def test_validate_without_runs(tmp_path, monkeypatch):
    monkeypatch.chdir(tmp_path)
    assert main(["validate"]) == 2
