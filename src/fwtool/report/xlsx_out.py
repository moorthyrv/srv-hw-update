"""report.xlsx: Summary, Worst 100, Model x Component matrix, Failures, Servers."""

from __future__ import annotations

import re
from collections import Counter, defaultdict
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from ..versions import version_key
from .rows import FIRMWARE_COLUMNS, SERVER_COLUMNS

HEADER_FILL = PatternFill("solid", fgColor="DDE3EA")
BOLD = Font(bold=True)
STATUS_FILLS = {
    "behind": PatternFill("solid", fgColor="F8D7DA"),
    "failed": PatternFill("solid", fgColor="E0E0E0"),
    "current": PatternFill("solid", fgColor="D4EDDA"),
    "partial": PatternFill("solid", fgColor="FFF3CD"),
}

WORST_COLUMNS = [
    "rank", "priority_score", "name", "bmc_ip", "vendor", "model", "generation", "support", "vxrail",
    "environment", "site", "overall_status", "components_behind", "components_noncompliant",
    "oldest_component_age_days", "oldest_component", "bios_version", "bios_latest", "bmc_version", "bmc_latest",
    "collection_status",
]
FAIL_COLUMNS = ["name", "bmc_ip", "vendor", "model", "bmc_type", "bmc_version", "collection_status",
                "error_class", "error", "fix_hint", "site", "environment"]


def _table(ws, columns: list[str], rows: list[dict], start_row: int = 1, status_col: str | None = None) -> int:
    for j, col in enumerate(columns, 1):
        cell = ws.cell(row=start_row, column=j, value=col)
        cell.font, cell.fill = BOLD, HEADER_FILL
    for i, row in enumerate(rows, start_row + 1):
        for j, col in enumerate(columns, 1):
            ws.cell(row=i, column=j, value=row.get(col))
        if status_col and row.get(status_col) in STATUS_FILLS:
            ws.cell(row=i, column=columns.index(status_col) + 1).fill = STATUS_FILLS[row[status_col]]
    return start_row + len(rows) + 1


def _finish(ws, columns: list[str], widths: dict[str, int] | None = None, filter_rows: int | None = None) -> None:
    ws.freeze_panes = "A2"
    for j, col in enumerate(columns, 1):
        w = (widths or {}).get(col) or min(max(len(col) + 2, 10), 40)
        ws.column_dimensions[get_column_letter(j)].width = w
    if filter_rows is not None and filter_rows > 0:
        ws.auto_filter.ref = f"A1:{get_column_letter(len(columns))}{filter_rows + 1}"


def _counts_block(ws, row: int, title: str, counter_rows: list[tuple], headers: list[str]) -> int:
    ws.cell(row=row, column=1, value=title).font = Font(bold=True, size=12)
    row += 1
    for j, h in enumerate(headers, 1):
        c = ws.cell(row=row, column=j, value=h)
        c.font, c.fill = BOLD, HEADER_FILL
    for vals in counter_rows:
        row += 1
        for j, v in enumerate(vals, 1):
            ws.cell(row=row, column=j, value=v)
    return row + 2


def write_xlsx(path: str | Path, server_rows: list[dict], firmware_rows: list[dict], meta: dict) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Summary"
    _summary(ws, server_rows, firmware_rows, meta)

    worst = [r for r in server_rows if r.get("rank")][:100]
    ws2 = wb.create_sheet("Worst 100")
    _table(ws2, WORST_COLUMNS, worst, status_col="overall_status")
    _finish(ws2, WORST_COLUMNS, {"name": 28, "model": 26, "oldest_component": 40}, len(worst))

    ws3 = wb.create_sheet("Model x Component")
    matrix = _matrix(firmware_rows)
    cols = ["vendor", "model", "category", "component", "servers", "distinct_versions", "latest_version",
            "servers_behind", "version_spread"]
    _table(ws3, cols, matrix)
    _finish(ws3, cols, {"model": 26, "component": 45, "version_spread": 90}, len(matrix))
    for row in ws3.iter_rows(min_row=2, min_col=9, max_col=9):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")

    fails = [r for r in server_rows if r.get("collection_status") in ("failed", "partial")]
    fails.sort(key=lambda r: (r["collection_status"] != "failed", r.get("error_class") or "", r["name"]))
    ws4 = wb.create_sheet("Failures & Partials")
    _table(ws4, FAIL_COLUMNS, fails, status_col="collection_status")
    _finish(ws4, FAIL_COLUMNS, {"error": 60, "fix_hint": 60, "name": 28}, len(fails))

    ws5 = wb.create_sheet("Servers")
    _table(ws5, SERVER_COLUMNS, server_rows, status_col="overall_status")
    _finish(ws5, SERVER_COLUMNS, {"name": 28, "model": 26, "error": 50, "fix_hint": 50}, len(server_rows))

    if len(firmware_rows) < 1_000_000:
        ws6 = wb.create_sheet("Firmware")
        _table(ws6, FIRMWARE_COLUMNS, firmware_rows, status_col="status")
        _finish(ws6, FIRMWARE_COLUMNS, {"component": 45, "name": 28, "model": 26, "note": 50}, len(firmware_rows))

    wb.save(path)


def _summary(ws, servers: list[dict], firmware: list[dict], meta: dict) -> None:
    ws["A1"] = "Firmware inventory summary"
    ws["A1"].font = Font(bold=True, size=14)
    r = 3
    for k, v in meta.items():
        ws.cell(row=r, column=1, value=k).font = BOLD
        ws.cell(row=r, column=2, value=str(v))
        r += 1
    r += 1

    def by(key_fn, rows):
        c = Counter(key_fn(x) for x in rows)
        return sorted(c.items(), key=lambda kv: (-kv[1], str(kv[0])))

    r = _counts_block(ws, r, "Servers by vendor", [(k or "unknown", v) for k, v in by(lambda x: x["vendor"], servers)],
                      ["vendor", "servers"])
    r = _counts_block(ws, r, "Collection status",
                      [(k, v) for k, v in by(lambda x: x["collection_status"], servers)], ["status", "servers"])
    r = _counts_block(ws, r, "Firmware status (overall, per server)",
                      [(k, v) for k, v in by(lambda x: x["overall_status"], servers)], ["status", "servers"])

    gen_rows = defaultdict(lambda: [0, 0])
    for s in servers:
        g = gen_rows[(s["vendor"] or "unknown", s["generation"] or "unknown", s["bmc_type"] or "")]
        g[0] += 1
        g[1] += s["overall_status"] == "behind"
    r = _counts_block(ws, r, "By generation",
                      [(k[0], k[1], k[2], v[0], v[1]) for k, v in sorted(gen_rows.items())],
                      ["vendor", "generation", "bmc", "servers", "behind"])

    model_rows = defaultdict(lambda: [0, 0, 0])
    for s in servers:
        m = model_rows[(s["vendor"] or "unknown", s["model"] or "unknown")]
        m[0] += 1
        m[1] += s["overall_status"] == "behind"
        m[2] += s["vxrail"] == "yes"
    r = _counts_block(ws, r, "By model",
                      [(k[0], k[1], v[0], v[1], v[2]) for k, v in sorted(model_rows.items(), key=lambda kv: -kv[1][0])],
                      ["vendor", "model", "servers", "behind", "vxrail"])

    r = _counts_block(ws, r, "Support type (HPE)",
                      [(k or "(not set)", v) for k, v in by(lambda x: x["support"], [s for s in servers if s["vendor"] == "hpe"])],
                      ["support", "servers"])

    cat = defaultdict(Counter)
    for f in firmware:
        cat[f["category"]][f["status"]] += 1
    statuses = ["current", "behind", "unknown", "no-reference"]
    r = _counts_block(ws, r, "Components by category and status",
                      [(k, *[cat[k][s] for s in statuses]) for k in sorted(cat)], ["category", *statuses])
    ws.column_dimensions["A"].width = 30
    ws.column_dimensions["B"].width = 28
    for col in "CDEF":
        ws.column_dimensions[col].width = 14


def _matrix(firmware: list[dict]) -> list[dict]:
    groups: dict[tuple, dict] = {}
    for f in firmware:
        key = (f["vendor"], f["model"], f["category"], _component_family(f["component"]))
        g = groups.setdefault(key, {"versions": Counter(), "servers": set(), "latest": "", "behind": set()})
        g["versions"][f["installed_version"]] += 1
        g["servers"].add(f["bmc_ip"])
        if f["latest_version"] and version_key(f["latest_version"]) > version_key(g["latest"]):
            g["latest"] = f["latest_version"]
        if f["status"] == "behind":
            g["behind"].add(f["bmc_ip"])
    out = []
    order = {c: i for i, c in enumerate(["BIOS", "BMC", "CPLD", "Storage", "NIC", "Drive", "PSU", "Other"])}
    for (vendor, model, category, comp), g in sorted(groups.items(), key=lambda kv: (kv[0][0], kv[0][1], order.get(kv[0][2], 9), kv[0][3])):
        spread = ", ".join(f"{v} ({n})" for v, n in sorted(g["versions"].items(), key=lambda kv: version_key(kv[0]), reverse=True))
        out.append({
            "vendor": vendor, "model": model, "category": category, "component": comp,
            "servers": len(g["servers"]), "distinct_versions": len(g["versions"]), "latest_version": g["latest"],
            "servers_behind": len(g["behind"]), "version_spread": spread,
        })
    return out


def _component_family(name: str) -> str:
    """Group per-slot entries: 'Disk 0 in Backplane 1 ...' / 'Power Supply 1' collapse together."""
    n = re.sub(r"\b(Disk|Drive|Bay|Slot|Port|PSU|Power Supply|Backplane)\s*\d+\b", r"\1 #", name or "")
    n = re.sub(r"\s+-\s+[0-9A-F:]{12,}$", "", n)  # NIC MAC suffixes
    return n.strip()
